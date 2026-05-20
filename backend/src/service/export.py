import csv
import io
import logging
from collections.abc import AsyncGenerator, Sequence
from typing import Any, Literal, TypedDict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from pydantic_core import ErrorDetails, InitErrorDetails

from schema.records import RecordData, RecordFull, Specimen

logger = logging.getLogger(__name__)

COLUMN_MAPPING: dict[str, str] = {
    "country": "Country",
    "region": "Region",
    "district": "District",
    "locality": "Locality",
    "is_manual_location": "Manual Location",
    "latitude": "Latitude",
    "longitude": "Longitude",
    "verbatimcoordinates": "Verbatim Coordinates",
    "coordinate_uncertainty": "Coordinate Uncertainty (m)",
    "georef_source": "Geo Referenced By",
    "location_remarks": "Location Remarks",
    "verbatim_date": "Date",
    "date_precision": "Date Precision",
    "is_interval": "Date Interval",
    "habitat": "Habitat",
    "sampling_protocol": "Sampling Protocol",
    "sampling_effort": "Sampling Effort",
    "sample_size_value": "Sample Size Value",
    "sample_size_unit": "Sample Size Unit",
    "event_remarks": "Event Remarks",
    "field_number": "Field Number",
    "catalog_number": "Catalog Number",
    "collection_code": "Collection Code",
    "recorded_by": "Recorded By",
    "family": "Family",
    "genus": "Genus",
    "species": "Species",
    "tax_verbatim": "Taxon Verbatim",
    "taxon_rank": "Taxon Rank",
    "type_status": "Type Status",
    "accepted_name": "Accepted Name",
    "taxon_remarks": "Taxon Remarks",
    "quantity_type": "Quantity Type",
    "occurrence_remarks": "Occurrence Remarks",
    "identification_remarks": "Identification Remarks",
}


SPECIMEN_HEADER_MAP: dict[
    str,
    tuple[
        Literal["male", "female", "none"],
        Literal["adult", "subadult", "juvenile", "none"],
    ],
] = {
    "Male Adult Quantity": ("male", "adult"),
    "Male Subadult Quantity": ("male", "subadult"),
    "Male Juvenile Quantity": ("male", "juvenile"),
    "Male Unknown Quantity": ("male", "none"),
    "Female Adult Quantity": ("female", "adult"),
    "Female Subadult Quantity": ("female", "subadult"),
    "Female Juvenile Quantity": ("female", "juvenile"),
    "Female Unknown Quantity": ("female", "none"),
    "Unknown Adult Quantity": ("none", "adult"),
    "Unknown Subadult Quantity": ("none", "subadult"),
    "Unknown Juvenile Quantity": ("none", "juvenile"),
    "Unknown Quantity": ("none", "none"),
}

REVERSE_COLUMN_MAPPING: dict[str, str] = {v: k for k, v in COLUMN_MAPPING.items()}


class SpecimenColumnError(TypedDict):
    sex: Literal["male", "female", "none"]
    life_stage: Literal["adult", "subadult", "juvenile", "none"]
    raw: str


ParseResult = tuple[RecordData | None, ValidationError | None]


async def _gen_none() -> AsyncGenerator[ParseResult]:
    """Empty generator — yields nothing."""
    if False:
        yield  # type: ignore


def is_row_empty(row: dict[str, Any | None]) -> bool:
    return all(v is None or str(v).strip() == "" for v in row.values())


def _error_details_to_init(e: ErrorDetails) -> InitErrorDetails:
    if "ctx" in e:
        return {
            "type": e["type"],
            "loc": e["loc"],
            "input": e["input"],
            "ctx": e["ctx"],
        }

    return {
        "type": e["type"],
        "loc": e["loc"],
        "input": e["input"],
    }


def _merge_errors(
    val: ValidationError | None,
    errors: list[SpecimenColumnError] | None,
) -> ValidationError:
    merged: list[InitErrorDetails] = []

    if val:
        merged = [_error_details_to_init(e) for e in val.errors()]

    if errors:
        merged.extend(_specimen_errors_to_details(errors))

    return ValidationError.from_exception_data("specimens", merged)


def _specimen_errors_to_details(
    errors: list[SpecimenColumnError],
) -> list[InitErrorDetails]:
    return [
        {
            "type": "value_error",
            "loc": ("specimens", err["sex"], err["life_stage"]),
            "input": err["raw"],
            "ctx": {"error": f"Invalid number: got '{err['raw']}'"},
        }
        for err in errors
    ]


def _parse_specimen_columns(
    row: dict[str, str | None],
) -> tuple[list[Specimen], list[SpecimenColumnError]]:
    specimens: list[Specimen] = []
    errors: list[SpecimenColumnError] = []
    for header, (sex_val, ls_val) in SPECIMEN_HEADER_MAP.items():
        raw = row.get(header)
        if raw is None:
            continue

        try:
            count = float(raw)
        except (ValueError, TypeError):
            errors.append({"sex": sex_val, "life_stage": ls_val, "raw": raw})
            continue

        if count > 0:
            specimens.append(
                Specimen(
                    sex=sex_val,
                    life_stage=ls_val,
                    count=count,
                )
            )
    return specimens, errors


def _specimens_to_12_columns(
    specimens: list[Specimen],
) -> dict[str, float]:
    result: dict[str, float] = dict.fromkeys(list(SPECIMEN_HEADER_MAP), 0.0)
    for sp in specimens:
        sex_label = "Unknown" if sp.sex == "none" else sp.sex.capitalize()
        ls_label = "Unknown" if sp.life_stage == "none" else sp.life_stage.capitalize()
        if sex_label == "Unknown" and ls_label == "Unknown":
            header = "Unknown Quantity"
        else:
            header = f"{sex_label} {ls_label} Quantity"
        result[header] += sp.count
    return result


def _row_to_record_data(row: dict[str, str | None]) -> ParseResult:
    """Convert a row dict to RecordData using pydantic validation."""
    data_dict: dict[str, Any] = {}
    for display_name, field in REVERSE_COLUMN_MAPPING.items():
        raw_value = row.get(display_name)
        if raw_value is not None:
            data_dict[field] = raw_value

    specimens, specimen_errors = _parse_specimen_columns(row)
    if specimens:
        data_dict["specimens"] = specimens

    try:
        record = RecordData.model_validate(data_dict)
    except ValidationError as e:
        bad_fields = {err["loc"][0] for err in e.errors()}
        cleaned = {k: v for k, v in data_dict.items() if k not in bad_fields}
        partial = RecordData.model_validate(cleaned)
        return (partial, _merge_errors(e, specimen_errors))

    if specimen_errors:
        return (record, _merge_errors(None, specimen_errors))

    return (record, None)


def _parse_excel_value(value: Any) -> Any:  # noqa: ANN401
    """Convert Excel boolean formulas and string representations to Python bool."""
    if value is None:
        return None
    if isinstance(value, str):
        upper = value.strip().upper()
        if upper in ("=TRUE()", "TRUE", "YES", "1"):
            return True
        if upper in ("=FALSE()", "FALSE", "NO", "0"):
            return False
    return value


async def read_excel(file_content: bytes) -> tuple[AsyncGenerator[ParseResult], int]:
    wb = load_workbook(filename=io.BytesIO(file_content))
    ws = wb.active

    if ws is None:
        wb.close()
        return (_gen_none(), 0)

    all_rows = list(ws.iter_rows(values_only=True))
    total = max(len(all_rows) - 1, 0)
    if total == 0 or len(all_rows) < 2:
        wb.close()
        return (_gen_none(), 0)

    headers = [str(cell) if cell is not None else "" for cell in all_rows[0]]
    data_rows = all_rows[1:]
    wb.close()

    async def _gen() -> AsyncGenerator[ParseResult]:
        for row in data_rows:
            row_dict = {
                headers[j]: _parse_excel_value(row[j])
                for j in range(len(headers))
                if j < len(row)
            }
            yield _row_to_record_data(row_dict)

    return (_gen(), total)


async def read_csv(file_content: bytes) -> tuple[AsyncGenerator[ParseResult], int]:
    text = file_content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    rows = [{k: (v if v else None) for k, v in row.items()} for row in reader]
    total = len(rows)

    async def _gen() -> AsyncGenerator[ParseResult]:
        for row_dict in rows:
            yield _row_to_record_data(row_dict)

    return (_gen(), total)


def records_to_excel(records: Sequence[RecordFull]) -> bytes:
    wb = Workbook()
    ws = wb.active

    if ws is None:
        logger.error("ws is None - workbook has no active sheet")
        raise RuntimeError("Workbook has no active sheet")

    headers = list(COLUMN_MAPPING.values()) + list(SPECIMEN_HEADER_MAP)
    ws.append(headers)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20
        ws.cell(row=1, column=col).font = Font(bold=True)

    for record in reversed(records):
        row = [getattr(record, field, None) for field in COLUMN_MAPPING]
        specimen_vals = _specimens_to_12_columns(record.specimens or [])
        row.extend(specimen_vals[h] for h in list(SPECIMEN_HEADER_MAP))
        ws.append(row)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def records_to_csv(records: Sequence[RecordFull]) -> str:
    output = io.StringIO()
    fieldnames = list(COLUMN_MAPPING.values()) + list(SPECIMEN_HEADER_MAP)
    writer = csv.DictWriter(output, fieldnames=fieldnames)
    writer.writeheader()

    for record in reversed(records):
        row: dict[str, object] = {
            COLUMN_MAPPING[field]: getattr(record, field, None)
            for field in COLUMN_MAPPING
        }
        specimen_vals = _specimens_to_12_columns(record.specimens or [])
        for h in list(SPECIMEN_HEADER_MAP):
            row[h] = specimen_vals[h]
        writer.writerow(row)

    return output.getvalue()
